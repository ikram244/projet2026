"""
api_knowledge_route.py
-----------------------
Blueprint Flask : gestion de la base de connaissance du chatbot (admin uniquement)

GET    /api/admin/knowledge          -> liste des entrees
POST   /api/admin/knowledge          -> creer une entree (texte manuel)
POST   /api/admin/knowledge/upload   -> creer une entree a partir d'un document (pdf/docx/xlsx)
PUT    /api/admin/knowledge/<id>     -> modifier une entree
DELETE /api/admin/knowledge/<id>     -> supprimer une entree
"""
from flask import Blueprint, request, jsonify, session

from db import (
    create_connaissance,
    get_connaissances,
    update_connaissance,
    delete_connaissance,
)

knowledge_bp = Blueprint("knowledge_bp", __name__)

ALLOWED_EXTENSIONS = {"pdf", "docx", "xlsx"}


def _extension_ok(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


def _extraire_texte_pdf(file_storage):
    from PyPDF2 import PdfReader
    reader = PdfReader(file_storage)
    texte = ""
    for page in reader.pages:
        texte += (page.extract_text() or "") + "\n"
    return texte.strip()


def _extraire_texte_docx(file_storage):
    import docx
    document = docx.Document(file_storage)
    return "\n".join(p.text for p in document.paragraphs).strip()


def _extraire_texte_xlsx(file_storage):
    import openpyxl
    wb = openpyxl.load_workbook(file_storage, data_only=True)
    lignes = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            valeurs = [str(v) for v in row if v is not None]
            if valeurs:
                lignes.append(" | ".join(valeurs))
    return "\n".join(lignes).strip()


@knowledge_bp.route("/api/admin/knowledge", methods=["GET"])
def list_knowledge():
    entries = get_connaissances()
    for e in entries or []:
        if e.get("date_creation"):
            e["date_creation"] = e["date_creation"].strftime("%Y-%m-%d %H:%M")
        if e.get("date_modification"):
            e["date_modification"] = e["date_modification"].strftime("%Y-%m-%d %H:%M")
    return jsonify(entries or [])


@knowledge_bp.route("/api/admin/knowledge", methods=["POST"])
def add_knowledge():
    data = request.get_json(force=True)

    categorie = (data.get("categorie") or "").strip() or None
    titre = (data.get("titre") or "").strip()
    contenu = (data.get("contenu") or "").strip()
    mots_cles = (data.get("mots_cles") or "").strip() or None

    utilisateur_id = session.get("user_id")
    if not utilisateur_id:
        return jsonify({"error": "utilisateur non connecte"}), 401
    if not titre or not contenu:
        return jsonify({"error": "Titre et contenu sont requis."}), 400

    new_id = create_connaissance(utilisateur_id, categorie, titre, contenu, mots_cles)
    if not new_id:
        return jsonify({"error": "Erreur lors de la creation."}), 500

    return jsonify({"id": new_id, "categorie": categorie, "titre": titre, "contenu": contenu, "mots_cles": mots_cles}), 201


@knowledge_bp.route("/api/admin/knowledge/upload", methods=["POST"])
def upload_knowledge():
    utilisateur_id = session.get("user_id")
    if not utilisateur_id:
        return jsonify({"error": "utilisateur non connecte"}), 401

    if "fichier" not in request.files:
        return jsonify({"error": "Aucun fichier envoye."}), 400

    fichier = request.files["fichier"]
    if fichier.filename == "" or not _extension_ok(fichier.filename):
        return jsonify({"error": "Format non supporte (pdf, docx, xlsx uniquement)."}), 400

    categorie = (request.form.get("categorie") or "").strip() or None
    titre = (request.form.get("titre") or fichier.filename).strip()

    ext = fichier.filename.rsplit(".", 1)[1].lower()
    try:
        if ext == "pdf":
            contenu = _extraire_texte_pdf(fichier)
        elif ext == "docx":
            contenu = _extraire_texte_docx(fichier)
        elif ext == "xlsx":
            contenu = _extraire_texte_xlsx(fichier)
        else:
            contenu = ""
    except Exception as e:
        return jsonify({"error": f"Erreur lors de la lecture du fichier : {e}"}), 400

    if not contenu:
        return jsonify({"error": "Impossible d'extraire du texte de ce fichier."}), 400

    new_id = create_connaissance(utilisateur_id, categorie, titre, contenu, None)
    if not new_id:
        return jsonify({"error": "Erreur lors de la creation."}), 500

    return jsonify({"id": new_id, "categorie": categorie, "titre": titre, "contenu": contenu}), 201


@knowledge_bp.route("/api/admin/knowledge/<int:knowledge_id>", methods=["PUT"])
def edit_knowledge(knowledge_id):
    data = request.get_json(force=True)

    categorie = (data.get("categorie") or "").strip() or None
    titre = (data.get("titre") or "").strip()
    contenu = (data.get("contenu") or "").strip()
    mots_cles = (data.get("mots_cles") or "").strip() or None

    if not titre or not contenu:
        return jsonify({"error": "Titre et contenu sont requis."}), 400

    updated = update_connaissance(knowledge_id, categorie, titre, contenu, mots_cles)
    if not updated:
        return jsonify({"error": "Entree introuvable ou erreur de mise a jour."}), 404

    return jsonify({"id": knowledge_id, "categorie": categorie, "titre": titre, "contenu": contenu, "mots_cles": mots_cles})


@knowledge_bp.route("/api/admin/knowledge/<int:knowledge_id>", methods=["DELETE"])
def remove_knowledge(knowledge_id):
    deleted = delete_connaissance(knowledge_id)
    if not deleted:
        return jsonify({"error": "Entree introuvable."}), 404
    return jsonify({"deleted": knowledge_id})